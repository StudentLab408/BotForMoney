import io
from dataclasses import dataclass
from decimal import Decimal

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Twips

from bot.templates.header_template import (
    A4_LONG_SIDE_TWIPS,
    A4_SHORT_SIDE_TWIPS,
    FONT_NAME,
    FONT_SIZE_PT,
    INSTITUTION_LINES,
    INTERNAL_COLUMNS,
    OFFICIAL_COLUMNS,
    PAGE_MARGIN_TWIPS,
    PURPOSE_TEMPLATE,
    TOTAL_LABEL,
)
from bot.utils.format import money
from bot.utils.time import month_name


@dataclass
class DocRow:
    index: int
    last_name: str
    first_name: str
    middle_name: str
    group_number: str
    amount: Decimal  # full amount, before the 25% lab share
    withheld: Decimal  # 25% lab share — shown only in the internal document
    basis_text: str


def _setup_page(doc: Document, landscape: bool) -> None:
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Twips(A4_LONG_SIDE_TWIPS), Twips(A4_SHORT_SIDE_TWIPS)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Twips(A4_SHORT_SIDE_TWIPS), Twips(A4_LONG_SIDE_TWIPS)
    margin = Twips(PAGE_MARGIN_TWIPS)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = margin

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(FONT_SIZE_PT)


def _fill_row(cells, values: list[str], widths: list[int], *, bold: bool = False, center: bool = False) -> None:
    for cell, value, width in zip(cells, values, widths, strict=True):
        cell.width = Twips(width)
        cell.text = value
        for paragraph in cell.paragraphs:
            if center:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = bold


def build_supplement_docx(year: int, month: int, rows: list[DocRow], *, official: bool) -> io.BytesIO:
    """official=True: clean list for the university. official=False: internal list with 25% share and basis."""
    columns = OFFICIAL_COLUMNS if official else INTERNAL_COLUMNS
    headers = [header for header, _ in columns]
    widths = [width for _, width in columns]

    doc = Document()
    _setup_page(doc, landscape=not official)

    for line in INSTITUTION_LINES:
        paragraph = doc.add_paragraph()
        paragraph.add_run(line).bold = True
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    purpose = doc.add_paragraph(PURPOSE_TEMPLATE.format(month_name=month_name(month), year=year))
    purpose.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _fill_row(table.rows[0].cells, headers, widths, bold=True, center=True)

    for row in rows:
        values = [str(row.index), row.last_name, row.first_name, row.middle_name, row.group_number, money(row.amount)]
        if not official:
            values += [money(row.withheld), row.basis_text]
        _fill_row(table.add_row().cells, values, widths)

    if not official:
        total_amount = sum((r.amount for r in rows), Decimal(0))
        total_withheld = sum((r.withheld for r in rows), Decimal(0))
        totals = ["", "", "", "", TOTAL_LABEL, money(total_amount), money(total_withheld), ""]
        _fill_row(table.add_row().cells, totals, widths, bold=True)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
