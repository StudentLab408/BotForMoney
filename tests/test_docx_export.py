from decimal import Decimal

from docx import Document

from bot.services.docx_export import DocRow, build_supplement_docx
from bot.templates.header_template import INTERNAL_COLUMNS, OFFICIAL_COLUMNS
from bot.utils.time import to_period

SEPTEMBER_2026 = to_period(2026, 9)


def _rows() -> list[DocRow]:
    return [
        DocRow(1, "Корецкий", "Илья", "Игоревич", "10706125", Decimal("200.00"), Decimal(50), "«Robo», лауреат"),
        DocRow(2, "Гайчук", "Фёдор", "Олегович", "10706125", Decimal("37.50"), Decimal(9), "Конференция «X»"),
    ]


def _table_text(doc) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in doc.tables[0].rows]


def test_official_docx_has_full_amount_and_no_internal_columns():
    doc = Document(build_supplement_docx(SEPTEMBER_2026, _rows(), official=True))
    table = _table_text(doc)

    assert table[0] == [header for header, _ in OFFICIAL_COLUMNS]
    assert table[1] == ["1", "Корецкий", "Илья", "Игоревич", "10706125", "200"]
    assert table[2][5] == "37,5"
    assert len(table) == 3  # header + rows, no totals

    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Белорусский национальный технический университет" in full_text
    assert "за сентябрь 2026 г." in full_text


def test_internal_docx_has_lab_share_basis_and_totals():
    doc = Document(build_supplement_docx(SEPTEMBER_2026, _rows(), official=False))
    table = _table_text(doc)

    assert table[0] == [header for header, _ in INTERNAL_COLUMNS]
    assert table[1][5:] == ["200", "50", "«Robo», лауреат"]
    assert table[-1] == ["", "", "", "", "Итого", "237,5", "59", ""]
    assert doc.sections[0].page_width > doc.sections[0].page_height  # landscape


def test_cell_widths_are_not_duplicated():
    doc = Document(build_supplement_docx(SEPTEMBER_2026, _rows(), official=False))
    cell = doc.tables[0].rows[1].cells[0]
    assert len(cell._tc.tcPr.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcW")) == 1


def test_empty_list():
    doc = Document(build_supplement_docx(to_period(2026, 1), [], official=True))
    assert len(doc.tables[0].rows) == 1
